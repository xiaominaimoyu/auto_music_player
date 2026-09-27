"""离线文本型 PDF/DOCX 导入。

本模块只负责“抽取文字 + 解析候选 + 产出可校对结果”，绝不直接写入曲库。
扫描 PDF、图片型 DOCX 和五线谱/简谱图片会明确报告需要 OMR，不会静默当作
空文本或错误简谱。
"""

from __future__ import annotations

import os
import re
import zipfile
from dataclasses import dataclass, field

from core.parser import parse_jianpu
from core.score_io import ImportResult
from core.score_model import MAX_BPM, MIN_BPM


MAX_DOCUMENT_BYTES = 500 * 1024 * 1024
MAX_PDF_PAGES = 200
MAX_TEXT_CHARS = 2_000_000
MAX_DOCX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
_BPM_RE = re.compile(r"(?:\bbpm\b|速度|节拍)\s*[:：=]?\s*(\d{2,3})", re.I)


class DocumentImportError(ValueError):
    """用户可理解的文档抽取或识别错误。"""


@dataclass(frozen=True)
class ExtractedDocument:
    path: str
    kind: str
    text: str
    pages: int = 0
    paragraphs: int = 0
    warnings: tuple[str, ...] = ()
    tool: str = ""


def _check_file(path: str) -> str:
    resolved = os.path.abspath(os.fspath(path))
    try:
        size = os.path.getsize(resolved)
    except OSError as exc:
        raise DocumentImportError(f"无法读取文档：{exc}") from exc
    if size > MAX_DOCUMENT_BYTES:
        raise DocumentImportError("文档不能超过 500 MB")
    return resolved


def _limit_text(text: str) -> str:
    text = str(text or "").replace("\x00", "")
    if len(text) > MAX_TEXT_CHARS:
        raise DocumentImportError("文档抽取文本超过 200 万字符")
    return text


def extract_pdf(path: str) -> ExtractedDocument:
    resolved = _check_file(path)
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentImportError(
            "缺少离线 PDF 文本抽取组件 pypdf；扫描 PDF 应转交 OMR 组件"
        ) from exc
    try:
        reader = PdfReader(resolved, strict=False)
        page_count = len(reader.pages)
        if page_count > MAX_PDF_PAGES:
            raise DocumentImportError(f"PDF 页数不能超过 {MAX_PDF_PAGES} 页")
        pages = []
        empty_pages = 0
        for index, page in enumerate(reader.pages):
            try:
                content = page.extract_text(extraction_mode="layout") or ""
            except TypeError:
                # 兼容较旧 pypdf 版本。
                content = page.extract_text() or ""
            content = _limit_text(content)
            if not content.strip():
                empty_pages += 1
            pages.append(f"[第 {index + 1} 页]\n{content}")
        warnings = []
        if empty_pages:
            warnings.append(
                f"有 {empty_pages} 页没有可抽取文字；若页面是扫描图片，请改用离线 OMR"
            )
        return ExtractedDocument(
            resolved,
            "pdf",
            _limit_text("\n\n".join(pages)),
            pages=page_count,
            warnings=tuple(warnings),
            tool="pypdf",
        )
    except DocumentImportError:
        raise
    except Exception as exc:
        raise DocumentImportError(f"PDF 文本抽取失败：{exc}") from exc


def _docx_text(document) -> tuple[str, int]:
    chunks = []
    count = 0
    if hasattr(document, "iter_inner_content"):
        blocks = document.iter_inner_content()
    else:
        blocks = list(document.paragraphs) + list(document.tables)
    for block in blocks:
        count += 1
        if hasattr(block, "text"):
            text = block.text
            if text and text.strip():
                chunks.append(text)
        else:
            for row in block.rows:
                cells = [cell.text for cell in row.cells]
                chunks.append("\t".join(cells))
    return "\n".join(chunks), count


def extract_docx(path: str) -> ExtractedDocument:
    resolved = _check_file(path)
    # 先检查 ZIP 解压总量，避免恶意 docx 在 python-docx 解析前耗尽内存。
    try:
        with zipfile.ZipFile(resolved) as archive:
            total = sum(max(0, int(info.file_size)) for info in archive.infolist())
            if total > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocumentImportError("DOCX 解压内容不能超过 200 MB")
    except DocumentImportError:
        raise
    except (OSError, zipfile.BadZipFile) as exc:
        raise DocumentImportError(f"DOCX 文件损坏：{exc}") from exc
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentImportError("缺少离线 DOCX 文本抽取组件 python-docx") from exc
    try:
        document = Document(resolved)
        text, paragraphs = _docx_text(document)
        warnings = []
        if not text.strip():
            warnings.append("DOCX 没有可抽取文字；嵌入图片需要离线 OMR")
        return ExtractedDocument(
            resolved,
            "docx",
            _limit_text(text),
            paragraphs=paragraphs,
            warnings=tuple(warnings),
            tool="python-docx",
        )
    except DocumentImportError:
        raise
    except Exception as exc:
        raise DocumentImportError(f"DOCX 文本抽取失败：{exc}") from exc


def _bpm_from_text(text: str) -> int:
    match = _BPM_RE.search(text or "")
    if not match:
        return 100
    value = int(match.group(1))
    return min(MAX_BPM, max(MIN_BPM, value))


def import_document(path: str) -> ImportResult:
    """抽取文本并生成候选乐谱；是否入库由 GUI 的确认门禁决定。"""

    extension = os.path.splitext(os.fspath(path))[1].lower()
    if extension == ".pdf":
        extracted = extract_pdf(path)
    elif extension == ".docx":
        extracted = extract_docx(path)
    else:
        raise DocumentImportError("仅支持文本型 .pdf 或 .docx；旧 .doc 需先另存为 .docx")

    notes, errors = parse_jianpu(extracted.text, collect=True, strict_ai=True)
    warnings = list(extracted.warnings)
    if not extracted.text.strip():
        warnings.append("没有抽取到可解析文本；请使用离线 OMR 组件")
    if errors:
        warnings.extend(
            f"第 {error.line} 行：{error.token}（{error.reason}）" for error in errors[:20]
        )
        if len(errors) > 20:
            warnings.append(f"另有 {len(errors) - 20} 个解析问题")
    metadata = {
        "schema_version": 1,
        "format": "document-text",
        "extension": extension,
        "tool": extracted.tool,
        "pages": extracted.pages,
        "paragraphs": extracted.paragraphs,
        "parse_error_count": len(errors),
        "manual_confirmation_required": True,
    }
    return ImportResult(
        name=os.path.splitext(os.path.basename(extracted.path))[0],
        bpm=_bpm_from_text(extracted.text),
        notes=notes,
        warnings=warnings,
        kind="document",
        source_format="document-text",
        raw_text=extracted.text,
        source_metadata=metadata,
    )


__all__ = [
    "DocumentImportError",
    "ExtractedDocument",
    "extract_docx",
    "extract_pdf",
    "import_document",
]
